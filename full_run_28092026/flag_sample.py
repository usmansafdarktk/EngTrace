"""A stratified sample of the digit rule's flags on this roster, for an author to read (D-150).

    python -m full_run_28092026.flag_sample --draw [--per-model 20] [--seed 0]   # FREE: the sample, local
    python -m full_run_28092026.flag_sample --reader-copy                        # FREE: the workbook a reader fills in
    python -m full_run_28092026.flag_sample --merge <returned .xlsx or .csv>     # FREE: its verdicts into sample.csv
    python -m full_run_28092026.flag_sample --score [--read-by "an author"]      # FREE: the verdicts, counts only

WHY. The digit rule's precision, three flags in four real, was measured on the pilot's 300 traces from
five frontier-plus-Llama models (SCORER_VALIDATION.md). This roster writes differently, and its flags
are uneven: Qwen3-235B-2507 has 1,547 on 18,785 claims, and 16 of DeepSeek's 24 sit in one template.
The pilot's rule for the checker was "validate on gold, then read every flag raised on real traces"
(`arith.py`), and reading is what found its last four parser defects. Nobody has read this roster's
flags. Q3 prints the digit rule's rates with the pilot's precision beside them; if this roster's is
lower, the caption overstates them.

WHAT --draw WRITES, under scores/flag_review/ (gitignored with the store; it holds trace text):
  sample.jsonl   one flagged claim per line: a code, the model, item and template ids, the step index,
                 the claim as the checker split it (left = right), the values it evaluated each side to,
                 the displayed precision that judged it, the unit tails, and the step's text
  sample.csv     the same claims, one per row, with empty `verdict` and `note` columns to fill in
Per model, up to --per-model flagged claims are taken round-robin across the model's templates, lowest
first within a template by a seeded shuffle, so no template dominates a model's sample. A claim is
drawn from the trace at the store's step index, recomputed with `arith.check`, so what the reader sees
is what the rule flagged.

HOW TO READ ONE. Recompute the left side from the numbers shown and ask whether the right side is a
correct rounding at the precision it displays. Verdicts:
  slip      the trace's arithmetic is wrong at the displayed digit: a real flag
  checker   the arithmetic is right and the checker misread it (a unit it did not know, a rounding
            chain it should tolerate, a clause split wrongly, a value inherited from the wrong side):
            a checker gap, to be fixed the D-137 way (measured on gold, pilot and full run first)
  unsure    the reader cannot tell
Say why in `note` for every `checker` verdict: the note is what the fix is built from.

THE READER'S COPY (--reader-copy). scores/flag_review/flag_review_reader.xlsx, for a reader outside the
code: one sheet of the claims in the sample's (shuffled) order with the step's full text and the units
the checker read added, the model's name left out so it cannot bias the reading, numbers kept as text
exactly as the checker saw them, and a slip / checker / unsure list in the verdict column. The
instructions are FLAG_READER_INSTRUCTIONS.md (committed: it holds no trace text), copied beside the
workbook as flag_review_instructions.md, so the two files to send sit together.
`--merge` takes the returned workbook (or a CSV, with or without a byte-order mark) and writes its
verdicts and notes into sample.csv by `code`, refusing a code or verdict it does not know. Like the
experts' labels, the filled files stay local and are never committed.

WHAT --score WRITES: FLAG_REVIEW.md, committed, counts only: per model the claims read, the verdicts,
precision among the decided ones (slip / (slip + checker)) with a Wilson interval, and the `checker`
notes grouped by their first word or two. No trace text. `--read-by` names who read them, for the title;
the paper must say the same. sample.csv is read with or without a byte-order mark.
"""
from __future__ import annotations

import argparse
import collections
import csv
import hashlib
import json
import math
import random
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
REPO = HERE.parent
for p in (str(REPO), str(REPO / 'evaluator_pilot_17092026' / 'evaluators'), str(REPO / 'evaluation')):
    if p not in sys.path:
        sys.path.insert(0, p)

import arith  # noqa: E402
import e2_prm  # noqa: E402

from full_run_28092026 import score  # noqa: E402
from full_run_28092026.analyze import ROSTER  # noqa: E402

OUT = score.SCORES / 'flag_review'
SAMPLE = OUT / 'sample.jsonl'
CSV = OUT / 'sample.csv'
READER = OUT / 'flag_review_reader.xlsx'
INSTRUCTIONS = HERE / 'FLAG_READER_INSTRUCTIONS.md'          # committed: what the reader is told
READER_NOTES = OUT / 'flag_review_instructions.md'           # the copy sent beside the workbook
VERDICTS = ('slip', 'checker', 'unsure')
FIELDS = ['code', 'model', 'template_id', 'item_id', 'step', 'claim', 'left_value', 'right_value',
          'displayed_ulp', 'verdict', 'note']


def code_of(model: str, item_id: str, step: int, k: int) -> str:
    return 'FL-' + hashlib.sha256(f'{model}|{item_id}|{step}|{k}'.encode('utf-8')).hexdigest()[:8]


def flagged_claims(model: str) -> list[dict]:
    """Every claim the digit rule flags in the model's answered traces, recomputed from the text."""
    rows = {r['item_id']: r for r in map(json.loads, (score.SCORES / 'main' / f'{model}.jsonl')
                                              .read_text(encoding='utf-8').splitlines())}
    texts = score.texts_matching('main', model, rows.values())
    out = []
    for item_id, r in rows.items():
        if r['status'] != 'answered' or not any(s['digit_flags'] for s in r['steps']):
            continue
        steps = e2_prm.steps_of(texts[item_id])
        if len(steps) != len(r['steps']):
            raise SystemExit(f'{model}/{item_id}: the trace splits into {len(steps)} steps, the store has {len(r["steps"])}')
        for j, (step, srow) in enumerate(zip(steps, r['steps'])):
            if not srow['digit_flags']:
                continue
            for k, c in enumerate(x for x in arith.check(step).claims if not x.ok_digit):
                out.append({'code': code_of(model, item_id, j, k), 'model': model, 'template_id': r['template_id'],
                            'item_id': item_id, 'step': j, 'line': c.line, 'left': c.left, 'right': c.right,
                            'left_value': c.left_value, 'right_value': c.right_value, 'displayed_ulp': c.ulp,
                            'left_unit': c.left_unit, 'right_unit': c.right_unit, 'ok_at_1pct': c.ok,
                            'step_text': step})
    return out


def draw(per_model: int, seed: int) -> int:
    rng = random.Random(seed)
    OUT.mkdir(parents=True, exist_ok=True)
    sample, totals = [], {}
    for model in ROSTER:
        claims = flagged_claims(model)
        totals[model] = len(claims)
        by_t = collections.defaultdict(list)
        for c in claims:
            by_t[c['template_id']].append(c)
        queues = [by_t[t] for t in sorted(by_t)]
        for q in queues:
            rng.shuffle(q)
        rng.shuffle(queues)
        chosen = []
        while len(chosen) < per_model and any(queues):
            for q in queues:
                if q and len(chosen) < per_model:
                    chosen.append(q.pop(0))
        sample += chosen
    rng.shuffle(sample)                                    # the reader sees models mixed, not in blocks
    with open(SAMPLE, 'w', encoding='utf-8', newline='\n') as fh:
        for c in sample:
            fh.write(json.dumps(c, ensure_ascii=False) + '\n')
    with open(CSV, 'w', encoding='utf-8', newline='') as fh:
        w = csv.DictWriter(fh, fieldnames=FIELDS)
        w.writeheader()
        for c in sample:
            w.writerow({'code': c['code'], 'model': c['model'], 'template_id': c['template_id'], 'item_id': c['item_id'],
                        'step': c['step'], 'claim': f"{c['left']} = {c['right']}",
                        'left_value': '; '.join(f'{v:.10g}' for v in c['left_value']),
                        'right_value': '; '.join(f'{v:.10g}' for v in c['right_value']),
                        'displayed_ulp': c['displayed_ulp'], 'verdict': '', 'note': ''})
    print(f'{len(sample)} flagged claims drawn ({per_model} per model at most, seed {seed}) from '
          + ', '.join(f'{m} {n}' for m, n in totals.items()))
    print(f'fill the verdict column of {CSV.relative_to(HERE)} (slip / checker / unsure), then --score')
    return 0


READER_COLUMNS = [('code', 13), ('problem', 30), ('claim', 38), ("checker's left value", 16), ('right value', 14),
                  ('displayed precision', 11), ('units (left | right)', 14), ('step', 90), ('verdict', 11),
                  ('note', 40)]


def reader_copy(path: Path = READER) -> int:
    """The workbook a reader fills in (the docstring's THE READER'S COPY), from sample.jsonl, in its order."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    import shutil
    claims = [json.loads(l) for l in SAMPLE.read_text(encoding='utf-8').splitlines()]
    top = Alignment(wrap_text=True, vertical='top')
    wb = Workbook()
    ws = wb.active
    ws.title = 'Claims'
    for j, (name, width) in enumerate(READER_COLUMNS, 1):
        cell = ws.cell(row=1, column=j, value=name)
        cell.font, cell.alignment = Font(bold=True), top
        ws.column_dimensions[get_column_letter(j)].width = width
    for i, c in enumerate(claims, 2):
        units = (' | '.join(u or '-' for u in (c['left_unit'], c['right_unit']))
                 if c['left_unit'] or c['right_unit'] else '')
        values = [c['code'], c['template_id'].removeprefix('template_'), f"{c['left']} = {c['right']}",
                  '; '.join(f'{v:.10g}' for v in c['left_value']), '; '.join(f'{v:.10g}' for v in c['right_value']),
                  str(c['displayed_ulp']), units, c['step_text'], '', '']
        for j, v in enumerate(values, 1):
            cell = ws.cell(row=i, column=j, value=v)
            cell.alignment, cell.number_format = top, '@'      # text: Excel keeps every digit as the checker saw it
    last = len(claims) + 1
    verdict_col = get_column_letter([n for n, _ in READER_COLUMNS].index('verdict') + 1)
    dv = DataValidation(type='list', formula1='"' + ','.join(VERDICTS) + '"', allow_blank=True,
                        showErrorMessage=True, errorTitle='Verdict', error='Choose slip, checker or unsure')
    ws.add_data_validation(dv)
    dv.add(f'{verdict_col}2:{verdict_col}{last}')
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(READER_COLUMNS))}{last}'
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    shutil.copyfile(INSTRUCTIONS, path.parent / READER_NOTES.name)
    print(f'{len(claims)} claims written to {path}; it holds trace text: send it privately, never commit it')
    print(f'the instructions to send with it: {path.parent / READER_NOTES.name} (a copy of {INSTRUCTIONS.name})')
    return 0


def merge(src: Path, csv_path: Path = CSV) -> int:
    """The returned workbook's or CSV's verdicts and notes into sample.csv, by code."""
    src = Path(src)
    if src.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        wb = load_workbook(src, read_only=True, data_only=True)
        try:
            ws = wb['Claims'] if 'Claims' in wb.sheetnames else wb.worksheets[0]
            rows = list(ws.iter_rows(values_only=True))
        finally:
            wb.close()                                     # read-only mode holds the file open until closed
        head = [str(h or '').strip().lower() for h in rows[0]]
        got = [dict(zip(head, ['' if v is None else str(v) for v in r])) for r in rows[1:]]
    else:
        got = [{(k or '').strip().lower(): v or '' for k, v in r.items()}
               for r in csv.DictReader(open(src, encoding='utf-8-sig', newline=''))]
    for need in ('code', 'verdict', 'note'):
        if not got or need not in got[0]:
            raise SystemExit(f'{src.name}: no "{need}" column')
    base = list(csv.DictReader(open(csv_path, encoding='utf-8-sig', newline='')))
    by_code = {r['code']: r for r in base}
    returned = {r['code'].strip(): r for r in got if r['code'].strip()}
    unknown = sorted(set(returned) - set(by_code))
    if unknown:
        raise SystemExit(f'{len(unknown)} codes in {src.name} are not in the sample: {unknown[:5]}')
    bad = sorted(c for c, r in returned.items() if r['verdict'].strip().lower() not in VERDICTS + ('',))
    if bad:
        raise SystemExit(f'verdicts must be one of {VERDICTS}: {bad[:5]}')
    changed = 0
    for c, r in returned.items():
        new = (r['verdict'].strip().lower(), r['note'].strip())
        old = (by_code[c]['verdict'], by_code[c]['note'])
        changed += old != ('', '') and old != new
        by_code[c]['verdict'], by_code[c]['note'] = new
    with open(csv_path, 'w', encoding='utf-8', newline='') as fh:
        w = csv.DictWriter(fh, fieldnames=FIELDS)
        w.writeheader()
        w.writerows(base)
    counts = collections.Counter(by_code[c]['verdict'] or 'blank' for c in returned)
    missing = len(by_code) - len(returned)
    no_note = sum(by_code[c]['verdict'] == 'checker' and not by_code[c]['note'] for c in returned)
    print(f'{len(returned)} rows merged into {csv_path.name}: ' + ', '.join(f'{k} {n}' for k, n in counts.most_common())
          + (f'; {missing} sample rows are not in the returned file' if missing else '')
          + (f'; {changed} rows had another verdict or note before' if changed else '')
          + (f'; {no_note} checker verdicts have no note' if no_note else ''))
    return 0


def wilson(k: int, n: int) -> tuple[float, float]:
    if not n:
        return float('nan'), float('nan')
    z, p = 1.96, k / n
    d = 1 + z * z / n
    c = (p + z * z / (2 * n)) / d
    h = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / d
    return c - h, c + h


def score_verdicts(csv_path: Path = CSV, out: Path | None = None, read_by: str = 'an author') -> int:
    rows = list(csv.DictReader(open(csv_path, encoding='utf-8-sig', newline='')))
    bad = [r['code'] for r in rows if r['verdict'].strip() and r['verdict'].strip().lower() not in VERDICTS]
    if bad:
        raise SystemExit(f'verdicts must be one of {VERDICTS}: {bad[:5]}')
    drawn = collections.Counter(r['model'] for r in rows)
    per = collections.defaultdict(collections.Counter)
    notes = collections.Counter()
    for r in rows:
        v = r['verdict'].strip().lower()
        if v:
            per[r['model']][v] += 1
        if v == 'checker' and r['note'].strip():
            notes[' '.join(r['note'].strip().lower().split()[:2])] += 1
    total = collections.Counter()
    L = [f'# The digit rule\'s flags on this roster, read by {read_by} (D-150)', '',
         'Generated by `flag_sample.py --score`; the sample and the verdicts are defined in its docstring. Counts '
         'only: the claims stay local under `scores/flag_review/`.', '',
         '| model | flags drawn | read | slip | checker | unsure | precision among decided | 95% Wilson |',
         '|---|---:|---:|---:|---:|---:|---:|---:|']
    for m in ROSTER:
        c = per[m]
        read = sum(c.values())
        decided = c['slip'] + c['checker']
        prec = c['slip'] / decided if decided else float('nan')
        lo, hi = wilson(c['slip'], decided)
        total.update(c)
        L.append(f"| `{m}` | {drawn[m]} | {read} | {c['slip']} | {c['checker']} | {c['unsure']} | "
                 f"{'' if not decided else f'{prec:.3f}'} | {'' if not decided else f'{lo:.3f} to {hi:.3f}'} |")
    decided = total['slip'] + total['checker']
    lo, hi = wilson(total['slip'], decided)
    prec_all = f"{total['slip'] / decided:.3f}" if decided else ''
    ci_all = f'{lo:.3f} to {hi:.3f}' if decided else ''
    L += [f"| all | {sum(drawn.values())} | {sum(total.values())} | {total['slip']} | {total['checker']} | {total['unsure']} | "
          f"{prec_all} | {ci_all} |", '',
          'The pilot measured 0.750 on the 300 labelled traces (SCORER_VALIDATION.md); a roster figure below it '
          'means Q3\'s digit-rule columns must carry this one instead.']
    if notes:
        L += ['', 'Checker verdicts by the note\'s first words: ' + ', '.join(f'{k} {n}' for k, n in notes.most_common()) + '.']
    (out or HERE / 'FLAG_REVIEW.md').write_text('\n'.join(L) + '\n', encoding='utf-8', newline='\n')
    print('\n'.join(L))
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split('\n')[0])
    ap.add_argument('--draw', action='store_true')
    ap.add_argument('--reader-copy', action='store_true')
    ap.add_argument('--merge', metavar='FILE', help='the returned workbook (.xlsx) or a CSV')
    ap.add_argument('--score', action='store_true')
    ap.add_argument('--read-by', default='an author', help='who read the flags, for the title of FLAG_REVIEW.md')
    ap.add_argument('--per-model', type=int, default=20)
    ap.add_argument('--seed', type=int, default=0)
    a = ap.parse_args()
    if a.draw:
        return draw(a.per_model, a.seed)
    if a.reader_copy:
        return reader_copy()
    if a.merge:
        return merge(Path(a.merge))
    if a.score:
        return score_verdicts(read_by=a.read_by)
    ap.print_help()
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
